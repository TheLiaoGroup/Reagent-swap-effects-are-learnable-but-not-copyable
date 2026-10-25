"""Export panel-specific Source Data workbooks from a figure manifest."""
from __future__ import annotations

import argparse
import hashlib
import json
import math
import os
from pathlib import Path
import re

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


def verify_workbooks(payload, work):
    """Read back every saved cell, worksheet name, table, and frozen pane."""
    work = Path(work)
    checks = []
    for book in payload["workbooks"]:
        book_path = Path(book["path"])
        if not book_path.is_absolute():
            book_path = work / book_path
        workbook = load_workbook(book_path, read_only=False, data_only=True)
        try:
            if workbook.sheetnames != [sheet["name"] for sheet in book["sheets"]]:
                raise ValueError(f"Worksheet mismatch: {book['path']}")
            total = 0
            for sheet in book["sheets"]:
                worksheet = workbook[sheet["name"]]
                expected = [sheet["columns"], *sheet["values"]]
                if (worksheet.max_row, worksheet.max_column) != (len(expected), len(sheet["columns"])):
                    raise ValueError(f"Worksheet dimensions mismatch: {worksheet.title}")
                for row_number, (actual, source) in enumerate(zip(worksheet.values, expected), 1):
                    for column_number, (observed, target) in enumerate(zip(actual, source), 1):
                        if target in (None, "") and observed in (None, ""):
                            continue
                        if isinstance(target, (float, int)) and not isinstance(target, bool):
                            equal = isinstance(observed, (float, int)) and not isinstance(observed, bool) and math.isclose(
                                observed, target, rel_tol=1e-12, abs_tol=0
                            )
                        else:
                            equal = observed == target and (not isinstance(target, bool) or isinstance(observed, bool))
                        if not equal:
                            raise ValueError(
                                f"Saved value mismatch: {worksheet.title}!R{row_number}C{column_number}: "
                                f"{observed!r} != {target!r}"
                            )
                        total += 1
                if worksheet.freeze_panes != "B2" or len(worksheet.tables) != 1:
                    raise ValueError(f"Missing frozen row/column or filter table: {worksheet.title}")
            checks.append({
                "path": book["path"],
                "sha256": hashlib.sha256(book_path.read_bytes()).hexdigest(),
                "sheets": len(book["sheets"]),
                "verified_cells": total,
                "exact_text_and_typed_numbers": True,
                "numeric_relative_tolerance": 1e-12,
                "freeze_panes": "B2",
                "filter_tables": True,
            })
        finally:
            workbook.close()
    (work / "source_data_validation.json").write_text(json.dumps(checks, indent=2), encoding="utf-8")
    print(
        f"Verified {sum(check['verified_cells'] for check in checks):,} saved source cells "
        f"in {sum(check['sheets'] for check in checks)} worksheets",
        flush=True,
    )


def _source_sheet(entry, manifest_directory):
    csv_path = Path(entry["csv"])
    if not csv_path.is_absolute():
        csv_path = manifest_directory / csv_path
    data = pd.read_csv(csv_path, keep_default_na=False)
    for column in entry.get("numeric_columns", []):
        data[column] = pd.to_numeric(data[column].replace("", None), errors="raise")
    if len(data) == 0:
        raise ValueError(f"Empty source table: {entry['name']}")
    return {
        "name": entry["name"][:31],
        "columns": list(data.columns),
        "values": data.astype(object).where(pd.notna(data), None).values.tolist(),
    }


def _save_workbook(book, work):
    workbook = Workbook()
    workbook.remove(workbook.active)
    try:
        for entry in book["sheets"]:
            worksheet = workbook.create_sheet(entry["name"])
            worksheet.sheet_view.showGridLines = False
            worksheet.freeze_panes = "B2"
            rows = [entry["columns"], *entry["values"]]
            for row_number, row in enumerate(rows, 1):
                for column_number, value in enumerate(row, 1):
                    cell = worksheet.cell(row_number, column_number, value)
                    if isinstance(value, str):
                        cell.data_type = "s"
                    cell.font = Font(name="Arial", size=10)
                    cell.alignment = Alignment(vertical="center", wrap_text=True)
                worksheet.row_dimensions[row_number].height = 18
            for cell in worksheet[1]:
                cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="34434B")
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            worksheet.row_dimensions[1].height = 38
            for column_number, name in enumerate(entry["columns"], 1):
                width = 52 if re.search(r"source|selector|definition|recipe|smiles|target_key|record_id", name) else 45 if "label" in name else 22
                worksheet.column_dimensions[get_column_letter(column_number)].width = width
                if re.search(r"value|yield|pp|mse|content|gain|contrast", name):
                    for row_number in range(2, len(rows) + 1):
                        worksheet.cell(row_number, column_number).number_format = "0.000000"
            table_name = "Data_" + book["id"] + "_" + re.sub(r"[^A-Za-z0-9_]", "_", entry["name"])
            table = Table(displayName=table_name, ref=f"A1:{get_column_letter(len(entry['columns']))}{len(rows)}")
            table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=False)
            worksheet.add_table(table)
        workbook.save(work / book["path"])
    finally:
        workbook.close()


def build_workbooks(manifest_path: Path, output: Path, work: Path):
    """Build the five workbooks, preserving manifest order within each book.

    Relative CSV paths are resolved against the manifest's directory.
    """
    manifest_path, output, work = Path(manifest_path), Path(output), Path(work)
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    output.mkdir(parents=True, exist_ok=True)
    work.mkdir(parents=True, exist_ok=True)
    selections = [
        (f"Fig{figure}", f"Source_Data_Fig{figure}.xlsx", f"Fig{figure}_")
        for figure in range(1, 5)
    ]
    selections.append(("SI", "Source_Data_Supplementary.xlsx", "S"))
    payload = {"workbooks": []}
    for name, filename, prefix in selections:
        sheets = [
            _source_sheet(entry, manifest_path.parent)
            for entry in manifest["source_data"] if entry["name"].startswith(prefix)
        ]
        if not sheets:
            raise ValueError(f"No source tables for {name}")
        book_path = Path(os.path.relpath(output / filename, start=work)).as_posix()
        payload["workbooks"].append({"id": name, "path": book_path, "sheets": sheets})
    (work / "workbook_payload.json").write_text(json.dumps(payload), encoding="utf-8")
    for book in payload["workbooks"]:
        _save_workbook(book, work)
    verify_workbooks(payload, work)
    checks = [{
        "path": book["path"],
        "sheets": [{"name": sheet["name"], "rows": len(sheet["values"]), "columns": len(sheet["columns"])} for sheet in book["sheets"]],
        "errorScan": "",
    } for book in payload["workbooks"]]
    (work / "source_data_checks.json").write_text(json.dumps(checks, indent=2), encoding="utf-8")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("manifest", type=Path)
    parser.add_argument("output", type=Path)
    parser.add_argument("work", type=Path)
    arguments = parser.parse_args()
    build_workbooks(arguments.manifest, arguments.output, arguments.work)
